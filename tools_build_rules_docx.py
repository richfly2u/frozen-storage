# -*- coding: utf-8 -*-
"""從 frozen-storage/index.html 的規則面板抽出文字，產生 Word (.docx)
   單一來源 = App 內的規則，避免兩邊不一致。
"""
import io, re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r'C:\Users\alan\frozen-storage\index.html'
OUT = r'C:\Users\alan\frozen-storage\冷凍庫使用規則.docx'

html = io.open(SRC, encoding='utf-8', errors='replace').read()

# --- 取出 <div class="rules"> ... 到 <div class="tiny" 之前 ---
i = html.find('<div class="rules">')
j = html.find('<div class="tiny"', i)
if i < 0 or j < 0:
    sys.exit('✗ 找不到規則區塊')
block = html[i:j]

def clean(t):
    t = re.sub(r'<br\s*/?>', '\n', t)
    t = re.sub(r'<[^>]+>', '', t)
    t = t.replace('&nbsp;', ' ').replace('&amp;', '&').replace('&lt;', '<').replace('&gt;', '>')
    return t.strip()

# --- 依原始順序抓出元素 ---
items = []
pattern = re.compile(
    r'<h3>(?P<h3>.*?)</h3>'
    r'|<div class="rn">(?P<rn>.*?)</div>'
    r'|<ul class="rul"[^>]*>(?P<ul>.*?)</ul>'
    r'|<div class="rtag">(?P<rtag>.*?)</div>'
    r'|<div class="rwarn">(?P<rwarn>.*?)</div>',
    re.S)
for m in pattern.finditer(block):
    d = {k: v for k, v in m.groupdict().items() if v is not None}
    key = list(d)[0]
    if key == 'ul':
        lis = [clean(x) for x in re.findall(r'<li>(.*?)</li>', d['ul'], re.S)]
        items.append(('ul', lis))
    else:
        items.append((key, clean(d[key])))

print('抽到元素 %d 個：' % len(items), [k for k, _ in items])

# --- 組 Word ---
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for attr, v in (('top_margin', 1.8), ('bottom_margin', 1.8), ('left_margin', 2.0), ('right_margin', 2.0)):
    setattr(sec, attr, Cm(v))

st = doc.styles['Normal']
st.font.name = 'Microsoft JhengHei'
st.font.size = Pt(11.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.4

def setfont(run, size=None, bold=None, color=None, name='Microsoft JhengHei'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)

def shade(p, hexcolor):
    e = OxmlElement('w:shd'); e.set(qn('w:val'), 'clear'); e.set(qn('w:fill'), hexcolor)
    p._p.get_or_add_pPr().append(e)

def border(p, hexcolor):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + side)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '12')
        el.set(qn('w:space'), '6'); el.set(qn('w:color'), hexcolor)
        b.append(el)
    pPr.append(b)

# 主標
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
setfont(p.add_run('冷凍庫使用規則'), 20, True, (0x0F, 0x3B, 0x57))
p.paragraph_format.space_after = Pt(6)

# --- QR（從 index.html 內嵌 SVG 解出的網址，用 segno 重繪成高解析 PNG）---
import segno
QR_URL = 'https://richfly2u.github.io/frozen-storage/#items'
QR_PNG = r'C:\Users\alan\frozen-storage\rules_qr.png'
segno.make(QR_URL, error='m').save(QR_PNG, scale=12, border=2)

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
p.add_run().add_picture(QR_PNG, width=Cm(3.0))
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
setfont(p.add_run('材料清單　掃碼開啟 App 材料清單'), 9.5, None, (0x66, 0x66, 0x66))

for kind, val in items:
    if kind == 'h3':
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
        setfont(p.add_run(val), 14, True, (0x0F, 0x3B, 0x57))
    elif kind == 'rn':
        for line in val.split('\n'):
            p = doc.add_paragraph()
            setfont(p.add_run(line), 11.5)
    elif kind == 'ul':
        for li in val:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Cm(0.9)
            setfont(p.add_run(li), 11)
    elif kind == 'rtag':
        p = doc.add_paragraph()
        border(p, '2F6FD0'); shade(p, 'F7FAFF')
        p.paragraph_format.space_before = Pt(10)
        for n, line in enumerate(val.split('\n')):
            if n: p.add_run().add_break()
            setfont(p.add_run(line), 11, None, (0x0B, 0x4B, 0x6B))
    elif kind == 'rwarn':
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        border(p, '8FBF87'); shade(p, 'F2FAF0')
        setfont(p.add_run(val), 12.5, True, (0x24, 0x5C, 0x2E))

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(14)
setfont(p.add_run('若有不完備之處，將再訂立修改　·　炊事組 2026/10/25 訂立'), 9.5, None, (0x66, 0x66, 0x66))

doc.save(OUT)
print('✓ 已產生:', OUT)
